"""All server-rendered pages behind the login."""
from __future__ import annotations

import secrets
from io import BytesIO
from datetime import date
from pathlib import Path

from fastapi import APIRouter, BackgroundTasks, Depends, File, Form, Request, UploadFile
from fastapi.responses import FileResponse, RedirectResponse, Response
from sqlalchemy import desc, func, select
from sqlalchemy.orm import Session, selectinload

from app.changelog import CHANGELOG
from app.config import HOURS_REPORT_RECIPIENT, PROJECT_UPLOADS_DIR
from app.db import get_db
from app.deps import require_admin, require_manager, require_user
from app.reporting import build_activity_pdf
from app.models import (
    Activity,
    Comment,
    Priority,
    Project,
    ProjectAttachment,
    ProgressReport,
    ProjectHealth,
    ProjectStatus,
    Task,
    TaskStatus,
    User,
    UserRole,
)
from app.notifications import send_hours_report_email, send_project_assignment_email
from app.security import hash_password, verify_password
from app.services import log_activity, recompute_health
from app.templating import templates

router = APIRouter()

MAX_PROJECT_FILE_BYTES = 25 * 1024 * 1024
PROJECT_FILE_TYPES = {
    ".png", ".jpg", ".jpeg", ".gif", ".webp", ".svg",
    ".pdf", ".doc", ".docx", ".xls", ".xlsx", ".ppt", ".pptx",
    ".txt", ".csv", ".zip",
}

AT_RISK_HEALTH = (ProjectHealth.at_risk, ProjectHealth.blocked, ProjectHealth.delayed)


def _projects(db: Session) -> list[Project]:
    return list(
        db.scalars(
            select(Project)
            .options(selectinload(Project.tasks), selectinload(Project.members))
            .order_by(Project.created_at.desc())
        )
    )


# ---------------------------------------------------------------------------
# Dashboard
# ---------------------------------------------------------------------------
@router.get("/dashboard")
def dashboard(request: Request, user: User = Depends(require_user), db: Session = Depends(get_db)):
    projects = _projects(db)
    active = [p for p in projects if p.is_active]
    at_risk = [p for p in active if p.health in AT_RISK_HEALTH]
    waiting = [p for p in projects if p.status == ProjectStatus.waiting_on_client]

    overall = round(sum(p.progress for p in active) / len(active)) if active else 0

    recent_done = list(
        db.scalars(
            select(Task)
            .options(selectinload(Task.project), selectinload(Task.assignees))
            .where(Task.status == TaskStatus.done)
            .order_by(Task.updated_at.desc())
            .limit(6)
        )
    )

    total_tasks = db.scalar(select(func.count(Task.id))) or 0
    done_tasks = db.scalar(select(func.count(Task.id)).where(Task.status == TaskStatus.done)) or 0
    blocked_tasks = db.scalar(select(func.count(Task.id)).where(Task.status == TaskStatus.blocked)) or 0
    # Count the working team (exclude the admin role).
    team_count = db.scalar(select(func.count(User.id)).where(User.role != UserRole.admin)) or 0

    return templates.TemplateResponse(request, "dashboard.html", {
        "user": user,
        "nav": "dashboard",
        "welcome": request.query_params.get("welcome") == "1",
        "active": sorted(active, key=lambda p: p.priority != Priority.critical),
        "at_risk": at_risk,
        "waiting": waiting,
        "overall": overall,
        "recent_done": recent_done,
        "stats": {
            "active": len(active),
            "total_projects": len(projects),
            "total_tasks": total_tasks,
            "done_tasks": done_tasks,
            "blocked_tasks": blocked_tasks,
            "at_risk": len(at_risk),
            "team": team_count,
        },
    })


# ---------------------------------------------------------------------------
# Projects
# ---------------------------------------------------------------------------
@router.get("/projects")
def projects_list(
    request: Request,
    status: str = "",
    health: str = "",
    user: User = Depends(require_user),
    db: Session = Depends(get_db),
):
    projects = _projects(db)
    selected_status = next((s for s in ProjectStatus if s.value == status), None)
    selected_health = next((h for h in ProjectHealth if h.value == health), None)
    if selected_status:
        projects = [p for p in projects if p.status == selected_status]
    if selected_health:
        projects = [p for p in projects if p.health == selected_health]
    planning_projects = [p for p in projects if p.status == ProjectStatus.planning]
    development_projects = [p for p in projects if p.status == ProjectStatus.development]
    testing_projects = [p for p in projects if p.status == ProjectStatus.testing]
    maintenance_projects = [p for p in projects if p.status == ProjectStatus.maintenance]
    other_projects = [p for p in projects if p.status not in {
        ProjectStatus.planning, ProjectStatus.development,
        ProjectStatus.testing, ProjectStatus.maintenance,
    }]
    return templates.TemplateResponse(request, "projects.html", {
        "user": user, "nav": "projects", "projects": projects,
        "planning_projects": planning_projects,
        "development_projects": development_projects,
        "testing_projects": testing_projects,
        "maintenance_projects": maintenance_projects,
        "other_projects": other_projects,
        "filters_active": bool(selected_status or selected_health),
    })


@router.get("/projects/new")
def new_project_form(request: Request, user: User = Depends(require_user), db: Session = Depends(get_db)):
    members = list(db.scalars(
        select(User).where(User.role == UserRole.developer).order_by(User.name)
    ))
    return templates.TemplateResponse(request, "project_new.html", {
        "user": user, "nav": "projects", "members": members,
        "statuses": list(ProjectStatus), "priorities": list(Priority),
    })


@router.post("/projects")
def create_project(
    background_tasks: BackgroundTasks,
    name: str = Form(...),
    client: str = Form(""),
    description: str = Form(""),
    status: str = Form("Planning"),
    priority: str = Form("Medium"),
    start_date: str = Form(""),
    due_date: str = Form(""),
    members: list[int] = Form(default=[]),
    user: User = Depends(require_user),
    db: Session = Depends(get_db),
):
    if not name.strip():
        return RedirectResponse("/projects/new", status_code=303)

    def parse(value: str):
        try:
            return date.fromisoformat(value) if value else None
        except ValueError:
            return None

    project = Project(
        name=name.strip(),
        client=client.strip(),
        description=description.strip(),
        status=next((s for s in ProjectStatus if s.value == status), ProjectStatus.planning),
        priority=next((p for p in Priority if p.value == priority), Priority.medium),
        start_date=parse(start_date),
        due_date=parse(due_date),
        created_by=user.id,
    )
    chosen = list(db.scalars(select(User).where(User.id.in_(members)))) if members else []
    # A developer who creates a project is automatically on it.
    if user.role == UserRole.developer and user not in chosen:
        chosen.append(user)
    project.members = chosen
    recompute_health(project)
    db.add(project)
    db.flush()
    log_activity(db, user=user, verb="created",
                 summary=f'created project "{project.name}"', project=project)
    db.commit()
    for member in chosen:
        background_tasks.add_task(
            send_project_assignment_email,
            recipient=member.email,
            recipient_name=member.name,
            project_id=project.id,
            project_name=project.name,
            assigned_by=user.name,
        )
    return RedirectResponse(f"/projects/{project.id}", status_code=303)


def _parse_date(value: str):
    try:
        return date.fromisoformat(value) if value else None
    except ValueError:
        return None


def _can_edit_project(user: User, project: Project) -> bool:
    """Admins/Managers may edit any project; the creator may edit their own."""
    return user.role in (UserRole.admin, UserRole.manager) or project.created_by == user.id


def _can_upload_project(user: User, project: Project) -> bool:
    return _can_edit_project(user, project) or user in project.members


def _project_file_path(attachment: ProjectAttachment) -> Path | None:
    base = PROJECT_UPLOADS_DIR.resolve()
    path = (base / attachment.filepath).resolve()
    if base not in path.parents or not path.is_file():
        return None
    return path


@router.post("/projects/{project_id}/attachments")
async def upload_project_attachment(
    project_id: int,
    file: UploadFile = File(...),
    user: User = Depends(require_user),
    db: Session = Depends(get_db),
):
    project = db.get(Project, project_id)
    if not project:
        return RedirectResponse("/projects", status_code=303)
    if not _can_upload_project(user, project):
        return RedirectResponse(f"/projects/{project_id}?upload_error=forbidden", status_code=303)

    filename = Path(file.filename or "").name
    suffix = Path(filename).suffix.lower()
    if not filename or suffix not in PROJECT_FILE_TYPES:
        return RedirectResponse(f"/projects/{project_id}?upload_error=type", status_code=303)

    contents = await file.read(MAX_PROJECT_FILE_BYTES + 1)
    if len(contents) > MAX_PROJECT_FILE_BYTES:
        return RedirectResponse(f"/projects/{project_id}?upload_error=size", status_code=303)

    stored_name = f"{secrets.token_hex(16)}{suffix}"
    project_dir = PROJECT_UPLOADS_DIR / str(project_id)
    project_dir.mkdir(parents=True, exist_ok=True)
    (project_dir / stored_name).write_bytes(contents)
    db.add(ProjectAttachment(
        project_id=project_id,
        filename=filename[:255],
        filepath=f"{project_id}/{stored_name}",
        content_type=file.content_type or "application/octet-stream",
        size=len(contents),
        uploaded_by=user.id,
    ))
    db.commit()
    return RedirectResponse(f"/projects/{project_id}?uploaded=1", status_code=303)


@router.get("/projects/{project_id}/attachments/{attachment_id}")
def download_project_attachment(
    project_id: int,
    attachment_id: int,
    user: User = Depends(require_user),
    db: Session = Depends(get_db),
):
    attachment = db.get(ProjectAttachment, attachment_id)
    if not attachment or attachment.project_id != project_id:
        return RedirectResponse(f"/projects/{project_id}", status_code=303)
    path = _project_file_path(attachment)
    if not path:
        return RedirectResponse(f"/projects/{project_id}?upload_error=missing", status_code=303)
    return FileResponse(path, media_type=attachment.content_type, filename=attachment.filename)


@router.get("/projects/{project_id}/edit")
def edit_project_form(project_id: int, request: Request,
                      user: User = Depends(require_user), db: Session = Depends(get_db)):
    project = db.get(Project, project_id)
    if not project:
        return RedirectResponse("/projects", status_code=303)
    if not _can_edit_project(user, project):
        return RedirectResponse(f"/projects/{project_id}", status_code=303)
    members = list(db.scalars(
        select(User).where(User.role == UserRole.developer).order_by(User.name)
    ))
    return templates.TemplateResponse(request, "project_edit.html", {
        "user": user, "nav": "projects", "project": project, "members": members,
        "statuses": list(ProjectStatus), "priorities": list(Priority),
        "healths": list(ProjectHealth),
    })


@router.post("/projects/{project_id}/edit")
def update_project(
    project_id: int,
    background_tasks: BackgroundTasks,
    name: str = Form(...),
    client: str = Form(""),
    description: str = Form(""),
    status: str = Form("Planning"),
    priority: str = Form("Medium"),
    health: str = Form(""),
    start_date: str = Form(""),
    due_date: str = Form(""),
    members: list[int] = Form(default=[]),
    user: User = Depends(require_user),
    db: Session = Depends(get_db),
):
    project = db.get(Project, project_id)
    if not project:
        return RedirectResponse("/projects", status_code=303)
    if not _can_edit_project(user, project):
        return RedirectResponse(f"/projects/{project_id}", status_code=303)
    if not name.strip():
        return RedirectResponse(f"/projects/{project_id}/edit", status_code=303)

    previous_member_ids = {member.id for member in project.members}
    project.name = name.strip()
    project.client = client.strip()
    project.description = description.strip()
    project.status = next((s for s in ProjectStatus if s.value == status), project.status)
    project.priority = next((p for p in Priority if p.value == priority), project.priority)
    if health:
        project.health = next((h for h in ProjectHealth if h.value == health), project.health)
    project.start_date = _parse_date(start_date)
    project.due_date = _parse_date(due_date)
    chosen = list(db.scalars(select(User).where(User.id.in_(members)))) if members else []
    project.members = chosen
    log_activity(db, user=user, verb="updated",
                 summary=f'updated project "{project.name}"', project=project)
    db.commit()
    for member in chosen:
        if member.id not in previous_member_ids:
            background_tasks.add_task(
                send_project_assignment_email,
                recipient=member.email,
                recipient_name=member.name,
                project_id=project.id,
                project_name=project.name,
                assigned_by=user.name,
            )
    return RedirectResponse(f"/projects/{project_id}", status_code=303)


@router.post("/projects/{project_id}/delete")
def delete_project(project_id: int, confirm: str = Form(""),
                   user: User = Depends(require_user), db: Session = Depends(get_db)):
    project = db.get(Project, project_id)
    if not project:
        return RedirectResponse("/projects", status_code=303)
    # Deleting a whole project is destructive -> Admins / Managers only.
    if user.role not in (UserRole.admin, UserRole.manager):
        return RedirectResponse(f"/projects/{project_id}?delete_error=forbidden", status_code=303)
    # Hard confirmation: the typed name must match exactly.
    if confirm.strip() != project.name:
        return RedirectResponse(f"/projects/{project_id}?delete_error=mismatch", status_code=303)

    name = project.name
    task_count = len(project.tasks)
    # Preserve the audit trail: keep historical activity rows, just unlink them
    # from the row that's about to disappear (the text summary keeps the context).
    for a in db.scalars(select(Activity).where(Activity.project_id == project_id)):
        a.project_id = None
        a.task_id = None
    for r in db.scalars(select(ProgressReport).where(ProgressReport.project_id == project_id)):
        r.project_id = None  # keep the developer's stand-up, just unlink the project
    for attachment in project.attachments:
        path = _project_file_path(attachment)
        if path:
            path.unlink(missing_ok=True)
    # Record the deletion itself (no project link, so it survives the delete).
    log_activity(db, user=user, verb="deleted",
                 summary=f'deleted project "{name}" ({task_count} task{"" if task_count == 1 else "s"})')
    db.delete(project)  # cascades tasks, comments, attachments and member links
    db.commit()
    return RedirectResponse("/projects", status_code=303)


@router.get("/projects/{project_id}")
def project_detail(project_id: int, request: Request,
                   user: User = Depends(require_user), db: Session = Depends(get_db)):
    project = db.get(Project, project_id)
    if not project:
        return RedirectResponse("/projects", status_code=303)
    columns = {status: [t for t in project.tasks if t.status == status] for status in TaskStatus}
    reports = list(
        db.scalars(
            select(ProgressReport)
            .options(selectinload(ProgressReport.user))
            .where(ProgressReport.project_id == project_id)
            .order_by(ProgressReport.created_at.desc())
            .limit(6)
        )
    )
    activities = list(
        db.scalars(
            select(Activity).options(selectinload(Activity.user))
            .where(Activity.project_id == project_id)
            .order_by(Activity.created_at.desc()).limit(10)
        )
    )
    return templates.TemplateResponse(request, "project_detail.html", {
        "user": user, "nav": "projects", "project": project,
        "columns": columns, "reports": reports, "activities": activities,
        "delete_error": request.query_params.get("delete_error"),
        "upload_error": request.query_params.get("upload_error"),
        "uploaded": request.query_params.get("uploaded") == "1",
        "can_delete": user.role in (UserRole.admin, UserRole.manager),
        "can_edit": _can_edit_project(user, project),
        "can_add": user.role in (UserRole.admin, UserRole.manager) or user in project.members,
        "can_upload": _can_upload_project(user, project),
    })


# ---------------------------------------------------------------------------
# Kanban board
# ---------------------------------------------------------------------------
@router.get("/board")
def board(request: Request, project_id: int | None = None,
          user: User = Depends(require_user), db: Session = Depends(get_db)):
    # Developers only see/select boards for their own projects; oversight sees all.
    if user.role == UserRole.developer:
        projects = sorted(user.projects, key=lambda p: p.name)
    else:
        projects = list(db.scalars(select(Project).order_by(Project.name)))

    if not projects:
        return templates.TemplateResponse(request, "board.html", {
            "user": user, "nav": "board", "projects": [], "project": None, "columns": {},
        })

    allowed_ids = {p.id for p in projects}
    project = db.get(Project, project_id) if project_id else None
    # Block access to a board the developer isn't on -> fall back to a default.
    if project is None or project.id not in allowed_ids:
        project = next((p for p in projects if p.is_active), projects[0])

    columns = {status: [t for t in project.tasks if t.status == status] for status in TaskStatus}
    return templates.TemplateResponse(request, "board.html", {
        "user": user, "nav": "board", "projects": projects,
        "project": project, "columns": columns,
    })


# ---------------------------------------------------------------------------
# Task detail (comments live here)
# ---------------------------------------------------------------------------
@router.get("/tasks/recent")
def recent_tasks(request: Request, user: User = Depends(require_user), db: Session = Depends(get_db)):
    tasks = list(
        db.scalars(
            select(Task)
            .options(selectinload(Task.project))
            .where(Task.status == TaskStatus.done)
            .order_by(Task.updated_at.desc())
            .limit(6)
        )
    )
    return templates.TemplateResponse(request, "tasks.html", {
        "user": user, "nav": "board", "tasks": tasks,
    })


@router.get("/tasks/{task_id}")
def task_detail(task_id: int, request: Request,
                user: User = Depends(require_user), db: Session = Depends(get_db)):
    task = db.get(Task, task_id)
    if not task:
        return RedirectResponse("/board", status_code=303)
    can_delete_task = user.role in (UserRole.admin, UserRole.manager) or user in task.assignees
    return templates.TemplateResponse(request, "task_detail.html", {
        "user": user, "nav": "board", "task": task, "members": task.project.members,
        "can_delete_task": can_delete_task,
    })


@router.post("/tasks/{task_id}/delete")
def delete_task(task_id: int, user: User = Depends(require_user), db: Session = Depends(get_db)):
    task = db.get(Task, task_id)
    if not task:
        return RedirectResponse("/board", status_code=303)
    # Developers may delete only tasks assigned to them; Admins/Managers, any.
    if user.role not in (UserRole.admin, UserRole.manager) and user not in task.assignees:
        return RedirectResponse(f"/tasks/{task_id}?error=forbidden", status_code=303)

    project = task.project
    title = task.title
    # Preserve the audit trail: unlink old activities, then record the deletion.
    for a in db.scalars(select(Activity).where(Activity.task_id == task_id)):
        a.task_id = None
    db.delete(task)  # cascades comments and attachments
    log_activity(db, user=user, verb="deleted",
                 summary=f'deleted task "{title}"', project=project)
    db.commit()
    return RedirectResponse(f"/board?project_id={project.id}", status_code=303)


@router.post("/tasks/{task_id}/comment")
def add_comment(task_id: int, body: str = Form(...),
                user: User = Depends(require_user), db: Session = Depends(get_db)):
    task = db.get(Task, task_id)
    if task and body.strip():
        db.add(Comment(task_id=task.id, user_id=user.id, body=body.strip()))
        log_activity(db, user=user, verb="commented",
                     summary=f'commented on "{task.title}"',
                     project=task.project, task=task)
        db.commit()
    return RedirectResponse(f"/tasks/{task_id}", status_code=303)


# ---------------------------------------------------------------------------
# Team view
# ---------------------------------------------------------------------------
@router.get("/team")
def team(request: Request, user: User = Depends(require_user), db: Session = Depends(get_db)):
    # Team View tracks who's doing the delivery work -> Developers only.
    # Managers/Admins are oversight roles (dashboard, activity log, reports).
    members = list(
        db.scalars(
            select(User)
            .options(selectinload(User.tasks))
            .where(User.role == UserRole.developer)
            .order_by(User.name)
        )
    )
    today = date.today()

    def board_for(member: User):
        active = [t for t in member.tasks if t.status not in (TaskStatus.done,)]
        working = [t for t in active if t.status == TaskStatus.in_progress]
        blocked = [t for t in member.tasks if t.status == TaskStatus.blocked]
        completed_today = [
            t for t in member.tasks
            if t.status == TaskStatus.done and t.updated_at and t.updated_at.date() == today
        ]
        return {
            "working": working or active[:3],
            "blocked": blocked,
            "completed_today": completed_today,
            "open_count": len(active),
        }

    cards = [(m, board_for(m)) for m in members]
    return templates.TemplateResponse(request, "team.html", {
        "user": user, "nav": "team", "cards": cards,
    })


# ---------------------------------------------------------------------------
# Daily progress reports (standup timeline)
# ---------------------------------------------------------------------------
@router.get("/progress")
def progress(
    request: Request,
    developer: str = "",
    date_from: str = "",
    date_to: str = "",
    user: User = Depends(require_user),
    db: Session = Depends(get_db),
):
    is_dev = user.role == UserRole.developer

    q = (
        select(ProgressReport)
        .options(selectinload(ProgressReport.user), selectinload(ProgressReport.project))
        .order_by(ProgressReport.report_date.desc(), ProgressReport.created_at.desc())
    )
    # Developers only ever see their own history; Admins/Managers see everyone.
    if is_dev:
        q = q.where(ProgressReport.user_id == user.id)
        dev_id = None
    else:
        dev_id = int(developer) if developer.isdigit() else None
        if dev_id:
            q = q.where(ProgressReport.user_id == dev_id)

    d_from = _parse_date(date_from)
    d_to = _parse_date(date_to)
    if d_from:
        q = q.where(ProgressReport.report_date >= d_from)
    if d_to:
        q = q.where(ProgressReport.report_date <= d_to)

    reports = list(db.scalars(q.limit(300)))
    # The submit form (developers only) should only offer the dev's own projects.
    projects = sorted(user.projects, key=lambda p: p.name)
    developers = list(db.scalars(
        select(User).where(User.role == UserRole.developer).order_by(User.name)
    ))
    # Today's reports keyed by project ("" = general) so the form can pre-fill
    # the right one when the developer switches projects.
    today_map = {}
    if is_dev:
        for r in db.scalars(select(ProgressReport).where(
            ProgressReport.user_id == user.id,
            ProgressReport.report_date == date.today(),
        )):
            key = str(r.project_id) if r.project_id else ""
            today_map[key] = {"today": r.today or "", "tomorrow": r.tomorrow or "",
                              "blocked": r.blocked or ""}

    return templates.TemplateResponse(request, "progress.html", {
        "user": user, "nav": "progress", "reports": reports,
        "projects": projects, "developers": developers, "today_map": today_map,
        "is_dev": is_dev,
        "f": {"developer": developer, "date_from": date_from, "date_to": date_to},
    })


@router.post("/progress")
def submit_progress(
    request: Request,
    today: str = Form(""),
    tomorrow: str = Form(""),
    blocked: str = Form(""),
    project_id: str = Form(""),
    user: User = Depends(require_user),
    db: Session = Depends(get_db),
):
    # Only developers file daily reports; oversight roles don't.
    if user.role != UserRole.developer:
        return RedirectResponse("/progress", status_code=303)
    pid = int(project_id) if project_id.isdigit() else None
    # One report per developer per project per day (a "general" report uses pid=None).
    existing = db.scalar(
        select(ProgressReport).where(
            ProgressReport.user_id == user.id,
            ProgressReport.report_date == date.today(),
            ProgressReport.project_id.is_(None) if pid is None else ProgressReport.project_id == pid,
        )
    )
    if existing:
        existing.today, existing.tomorrow, existing.blocked = today, tomorrow, blocked
    else:
        db.add(ProgressReport(user_id=user.id, project_id=pid, report_date=date.today(),
                              today=today, tomorrow=tomorrow, blocked=blocked))
        log_activity(db, user=user, verb="reported", summary="submitted a daily progress report")
    db.commit()
    return RedirectResponse("/progress", status_code=303)


# ---------------------------------------------------------------------------
# Monthly hours tracker
# ---------------------------------------------------------------------------
def _hours_template(request: Request, user: User, customers: list[str], **values):
    return templates.TemplateResponse(request, "hours_tracker.html", {
        "user": user,
        "nav": "hours",
        "customers": customers,
        "month_default": date.today().strftime("%Y-%m"),
        **values,
    })


@router.get("/hours-tracker")
def hours_tracker(request: Request, user: User = Depends(require_user), db: Session = Depends(get_db)):
    customers = sorted({
        client.strip() for client in db.scalars(select(Project.client))
        if client and client.strip()
    }, key=str.casefold)
    return _hours_template(request, user, customers,
                           submitted=request.query_params.get("submitted") == "1")


@router.post("/hours-tracker")
def submit_hours(
    background_tasks: BackgroundTasks,
    request: Request,
    customer: str = Form(...),
    other_customer: str = Form(""),
    duration: str = Form(...),
    month: str = Form(...),
    description: str = Form(""),
    user: User = Depends(require_user),
    db: Session = Depends(get_db),
):
    customers = sorted({
        client.strip() for client in db.scalars(select(Project.client))
        if client and client.strip()
    }, key=str.casefold)
    customer = customer.strip()
    other_customer = other_customer.strip()
    description = description.strip()
    errors = []
    if customer not in customers and customer != "Other":
        errors.append("Select a customer from the list.")
    if customer == "Other" and not other_customer:
        errors.append("Enter the customer name when selecting Other.")
    try:
        hours = float(duration)
        if hours <= 0 or hours > 744:
            raise ValueError
    except ValueError:
        errors.append("Duration must be a number greater than 0 and no more than 744.")
    try:
        parsed_month = date.fromisoformat(f"{month}-01")
    except ValueError:
        parsed_month = None
        errors.append("Select a valid month.")
    if not HOURS_REPORT_RECIPIENT:
        errors.append("Monthly hours reporting email is not configured yet.")
    if errors:
        return _hours_template(request, user, customers, error=" ".join(errors),
                               values={"customer": customer, "other_customer": other_customer,
                                       "duration": duration, "month": month,
                                       "description": description})

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        return _hours_template(
            request,
            user,
            customers,
            error="Excel export is unavailable because the openpyxl package is not installed.",
            values={"customer": customer, "other_customer": other_customer,
                    "duration": duration, "month": month, "description": description},
        )

    customer_name = other_customer if customer == "Other" else customer
    month_label = parsed_month.strftime("%B %Y")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Monthly Hours"
    headers = ["Internal customer", "Duration (hours)", "Month", "Description", "IT Tech"]
    sheet.append(headers)
    sheet.append([customer_name, hours, month_label, description, user.name])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17324D")
    for column, width in zip("ABCDE", (28, 18, 18, 55, 24)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    output = BytesIO()
    workbook.save(output)
    filename = f"hours-{parsed_month.strftime('%Y-%m')}-{user.name.replace(' ', '-')}.xlsx"
    background_tasks.add_task(
        send_hours_report_email,
        recipient=HOURS_REPORT_RECIPIENT,
        submitter_email=user.email,
        submitter_name=user.name,
        month=month_label,
        customer=customer_name,
        duration=f"{hours:g}",
        description=description,
        workbook=output.getvalue(),
        filename=filename,
    )
    return RedirectResponse("/hours-tracker?submitted=1", status_code=303)


# ---------------------------------------------------------------------------
# Reports & analytics (the "weekly report" + charts)
# ---------------------------------------------------------------------------
@router.get("/reports")
def reports(request: Request, user: User = Depends(require_manager), db: Session = Depends(get_db)):
    projects = _projects(db)
    members = list(db.scalars(select(User).options(selectinload(User.tasks))))

    status_breakdown = {s: 0 for s in TaskStatus}
    for t in db.scalars(select(Task)):
        status_breakdown[t.status] += 1

    by_developer = sorted(
        [(m, sum(1 for t in m.tasks if t.status == TaskStatus.done), len(m.tasks)) for m in members],
        key=lambda x: x[1], reverse=True,
    )
    projects_by_status = {s: [p for p in projects if p.status == s] for s in ProjectStatus}

    return templates.TemplateResponse(request, "reports.html", {
        "user": user, "nav": "reports", "projects": projects,
        "status_breakdown": status_breakdown, "by_developer": by_developer,
        "projects_by_status": projects_by_status,
        "max_dev": max((d[1] for d in by_developer), default=1) or 1,
    })


# ---------------------------------------------------------------------------
# Admin: team member management
# ---------------------------------------------------------------------------
@router.get("/admin/users")
def admin_users(request: Request, user: User = Depends(require_admin), db: Session = Depends(get_db)):
    members = list(db.scalars(select(User).order_by(User.created_at)))
    return templates.TemplateResponse(request, "admin_users.html", {
        "user": user, "nav": "admin", "members": members, "roles": list(UserRole),
        "created": request.query_params.get("created") == "1",
        "updated": request.query_params.get("updated") == "1",
        "error": request.query_params.get("error"),
    })


@router.post("/admin/users")
def create_user(
    name: str = Form(...),
    email: str = Form(...),
    password: str = Form(...),
    role: str = Form("Developer"),
    title: str = Form("Developer"),
    accent: str = Form("#1CC4D8"),
    user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    email = email.strip().lower()
    if db.scalar(select(User).where(User.email == email)):
        return RedirectResponse("/admin/users?error=exists", status_code=303)
    role_enum = next((r for r in UserRole if r.value == role), UserRole.developer)
    new_user = User(name=name.strip(), email=email, password_hash=hash_password(password),
                    role=role_enum, title=title.strip() or role_enum.value, accent=accent)
    db.add(new_user)
    log_activity(db, user=user, verb="created",
                 summary=f'added team member {new_user.name} ({role_enum.value})')
    db.commit()
    return RedirectResponse("/admin/users?created=1", status_code=303)


@router.post("/admin/users/{user_id}")
def update_user(
    user_id: int,
    name: str = Form(...),
    email: str = Form(...),
    role: str = Form("Developer"),
    title: str = Form(""),
    accent: str = Form("#1CC4D8"),
    password: str = Form(""),
    user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    target = db.get(User, user_id)
    if not target:
        return RedirectResponse("/admin/users?error=missing", status_code=303)
    email = email.strip().lower()
    clash = db.scalar(select(User).where(User.email == email, User.id != user_id))
    if clash:
        return RedirectResponse("/admin/users?error=exists", status_code=303)
    target.name = name.strip()
    target.email = email
    target.role = next((r for r in UserRole if r.value == role), target.role)
    target.title = title.strip() or target.role.value
    target.accent = accent
    pw_note = ""
    if password.strip():
        target.password_hash = hash_password(password.strip())
        pw_note = " (password reset)"
    log_activity(db, user=user, verb="updated",
                 summary=f'updated team member {target.name}{pw_note}')
    db.commit()
    return RedirectResponse("/admin/users?updated=1", status_code=303)


# ---------------------------------------------------------------------------
# Account -- self-service password change (any signed-in user)
# ---------------------------------------------------------------------------
@router.get("/account")
def account(request: Request, user: User = Depends(require_user)):
    return templates.TemplateResponse(request, "account.html", {
        "user": user, "nav": "account",
        "updated": request.query_params.get("pw") == "1",
        "color_saved": request.query_params.get("color") == "1",
        "error": request.query_params.get("error"),
        "palette": ACCENT_PALETTE,
    })


@router.post("/account/password")
def change_password(
    current: str = Form(...),
    new: str = Form(...),
    confirm: str = Form(...),
    user: User = Depends(require_user),
    db: Session = Depends(get_db),
):
    if not verify_password(current, user.password_hash):
        return RedirectResponse("/account?error=wrong", status_code=303)
    if len(new) < 6:
        return RedirectResponse("/account?error=short", status_code=303)
    if new != confirm:
        return RedirectResponse("/account?error=mismatch", status_code=303)
    user.password_hash = hash_password(new)
    db.commit()
    return RedirectResponse("/account?pw=1", status_code=303)


# Preset avatar colours offered on the Account page (any hex is still allowed).
ACCENT_PALETTE = ["#1CC4D8", "#3FE0C2", "#E0B23C", "#7C9CF5", "#B08CF0",
                  "#F4677A", "#F0944C", "#6FB1BF", "#AEB9C7"]


def _valid_hex(value: str) -> bool:
    value = value.strip()
    if len(value) != 7 or value[0] != "#":
        return False
    try:
        int(value[1:], 16)
        return True
    except ValueError:
        return False


@router.post("/account/color")
def change_color(
    accent: str = Form(...),
    user: User = Depends(require_user),
    db: Session = Depends(get_db),
):
    """Any signed-in user may change their own avatar colour."""
    if not _valid_hex(accent):
        return RedirectResponse("/account?error=color", status_code=303)
    user.accent = accent.strip().upper()
    db.commit()
    return RedirectResponse("/account?color=1", status_code=303)


# ---------------------------------------------------------------------------
# What's new -- patch notes / changelog (any signed-in user)
# ---------------------------------------------------------------------------
@router.get("/changelog")
def changelog_page(request: Request, user: User = Depends(require_user)):
    return templates.TemplateResponse(request, "changelog.html", {
        "user": user, "nav": "changelog", "entries": CHANGELOG,
    })


# ---------------------------------------------------------------------------
# Activity log -- permanent audit trail (Admins / Managers only)
# ---------------------------------------------------------------------------
ACTIVITY_VERBS = ["created", "updated", "deleted", "moved", "completed",
                  "commented", "reported", "blocked"]


def _int_or_none(value: str | None):
    return int(value) if value and value.isdigit() else None


def _activity_query(db, *, developer, project, verb, date_from, date_to):
    q = (
        select(Activity)
        .options(selectinload(Activity.user), selectinload(Activity.project))
    )
    if developer:
        q = q.where(Activity.user_id == developer)
    if project:
        q = q.where(Activity.project_id == project)
    if verb:
        q = q.where(Activity.verb == verb)
    if date_from:
        q = q.where(func.date(Activity.created_at) >= date_from)
    if date_to:
        q = q.where(func.date(Activity.created_at) <= date_to)
    return q.order_by(Activity.created_at.desc())


def _filters_desc(*, users, projects, developer, project, verb, date_from, date_to) -> str:
    parts = []
    if developer:
        u = next((u for u in users if u.id == developer), None)
        parts.append(f"developer = {u.name if u else developer}")
    if project:
        p = next((p for p in projects if p.id == project), None)
        parts.append(f"project = {p.name if p else project}")
    if verb:
        parts.append(f"type = {verb}")
    if date_from:
        parts.append(f"from {date_from}")
    if date_to:
        parts.append(f"to {date_to}")
    return ", ".join(parts) if parts else "all activity, all time"


@router.get("/activity")
def activity_log(
    request: Request,
    page: int = 1,
    developer: str = "",
    project: str = "",
    verb: str = "",
    date_from: str = "",
    date_to: str = "",
    user: User = Depends(require_manager),
    db: Session = Depends(get_db),
):
    dev_id, proj_id = _int_or_none(developer), _int_or_none(project)
    verb = verb if verb in ACTIVITY_VERBS else ""
    q = _activity_query(db, developer=dev_id, project=proj_id, verb=verb,
                        date_from=date_from, date_to=date_to)

    total = db.scalar(select(func.count()).select_from(q.subquery())) or 0
    per_page = 60
    page = max(1, page)
    pages = max(1, (total + per_page - 1) // per_page)
    items = list(db.scalars(q.offset((page - 1) * per_page).limit(per_page)))

    users = list(db.scalars(select(User).order_by(User.name)))
    projects = list(db.scalars(select(Project).order_by(Project.name)))

    return templates.TemplateResponse(request, "activity.html", {
        "user": user, "nav": "activity", "items": items,
        "total": total, "page": page, "pages": pages,
        "users": users, "projects": projects, "verbs": ACTIVITY_VERBS,
        "f": {"developer": developer, "project": project, "verb": verb,
              "date_from": date_from, "date_to": date_to},
    })


@router.get("/activity/export.pdf")
def activity_pdf(
    developer: str = "",
    project: str = "",
    verb: str = "",
    date_from: str = "",
    date_to: str = "",
    user: User = Depends(require_manager),
    db: Session = Depends(get_db),
):
    dev_id, proj_id = _int_or_none(developer), _int_or_none(project)
    verb = verb if verb in ACTIVITY_VERBS else ""
    q = _activity_query(db, developer=dev_id, project=proj_id, verb=verb,
                        date_from=date_from, date_to=date_to)
    items = list(db.scalars(q))  # all matching rows, all time

    users = list(db.scalars(select(User)))
    projects = list(db.scalars(select(Project)))
    desc_text = _filters_desc(users=users, projects=projects, developer=dev_id,
                              project=proj_id, verb=verb, date_from=date_from, date_to=date_to)

    pdf = build_activity_pdf(items, filters_desc=desc_text, generated_by=user.name)
    filename = f"keystone-activity-{date.today().isoformat()}.pdf"
    return Response(
        content=pdf,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
