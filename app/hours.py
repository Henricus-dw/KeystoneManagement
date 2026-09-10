"""South African working-day rules for monthly hours reporting."""
from __future__ import annotations

from calendar import monthrange
from datetime import date, timedelta
from io import BytesIO
import json
from pathlib import Path

import holidays


SAST = "Africa/Johannesburg"
REQUIRED_MEMBER_NAMES = {"henri", "darryl", "jean", "taskeen"}


def member_key(user) -> str:
    return user.name.strip().split()[0].casefold() if user.name.strip() else ""


def is_required_member(user) -> bool:
    return member_key(user) in REQUIRED_MEMBER_NAMES


def south_african_holidays(year: int):
    return holidays.SouthAfrica(years=[year])


def is_working_day(day: date) -> bool:
    return day.weekday() < 5 and day not in south_african_holidays(day.year)


def first_working_day(day: date) -> date:
    while not is_working_day(day):
        day += timedelta(days=1)
    return day


def report_due_date(report_month: date) -> date:
    next_month = (report_month.replace(day=28) + timedelta(days=4)).replace(day=1)
    return first_working_day(next_month)


def working_hours(report_month: date) -> int:
    total = 0
    for day_number in range(1, monthrange(report_month.year, report_month.month)[1] + 1):
        if is_working_day(date(report_month.year, report_month.month, day_number)):
            total += 9
    return total


def previous_month(day: date) -> date:
    first = day.replace(day=1)
    return (first - timedelta(days=1)).replace(day=1)


def build_consolidated_workbook(report_month: date, submissions) -> bytes:
    """Build one worksheet per submitted team member."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    workbook = Workbook()
    workbook.remove(workbook.active)
    headers = ["Internal customer", "Duration (hours)", "Description", "Month", "Year", "IT Tech"]
    for submission in sorted(submissions, key=lambda item: item.user.name.casefold()):
        sheet = workbook.create_sheet(title=submission.user.name[:31])
        sheet.append(headers)
        for entry in json.loads(submission.entries):
            customer = entry["other_customer"] if entry["customer"] == "Other" else entry["customer"]
            sheet.append([
                customer,
                float(entry["duration"]),
                entry["description"],
                report_month.strftime("%B"),
                report_month.year,
                submission.user.name,
            ])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17324D")
        for column, width in zip("ABCDEF", (28, 18, 55, 18, 12, 24)):
            sheet.column_dimensions[column].width = width
        sheet.freeze_panes = "A2"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def save_workbook(path: Path, workbook: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(workbook)


def workbook_sheets(path: Path) -> list[dict]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=False)
    sheets = []
    for sheet in workbook.worksheets:
        sheets.append({
            "title": sheet.title,
            "rows": [[cell.value if cell.value is not None else "" for cell in row]
                     for row in sheet.iter_rows()],
        })
    return sheets


def update_workbook(path: Path, sheets: list[dict]) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    for sheet_data in sheets:
        sheet = workbook[sheet_data["title"]]
        for row_index, values in enumerate(sheet_data["rows"], start=1):
            for column_index, value in enumerate(values, start=1):
                sheet.cell(row=row_index, column=column_index).value = value
    workbook.save(path)


def update_member_workbook(path: Path, report_month: date, submission) -> None:
    """Replace only one member's worksheet after that member resubmits."""
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill

    workbook = load_workbook(path)
    title = submission.user.name[:31]
    if title in workbook.sheetnames:
        del workbook[title]
    sheet = workbook.create_sheet(title=title)
    sheet.append(["Internal customer", "Duration (hours)", "Description", "Month", "Year", "IT Tech"])
    for entry in json.loads(submission.entries):
        customer = entry["other_customer"] if entry["customer"] == "Other" else entry["customer"]
        sheet.append([
            customer,
            float(entry["duration"]),
            entry["description"],
            report_month.strftime("%B"),
            report_month.year,
            submission.user.name,
        ])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17324D")
    for column, width in zip("ABCDEF", (28, 18, 55, 18, 12, 24)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    workbook.save(path)